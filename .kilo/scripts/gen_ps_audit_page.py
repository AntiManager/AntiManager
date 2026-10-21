"""Generate the /genome/audit/ page regions and the fillable XLSX from the vault.

Source of truth: the Obsidian vault note «Аудит ПС/Экспресс-чек-лист-аудита-ПС.md».
The script fills the <!--AUDIT:*:START-->…<!--AUDIT:*:END--> marker regions in
site/src/content/genome/audit.html in place (authored prose outside the markers is
never touched) and writes site/materials/ps-express-audit/express-audit-ps.xlsx.

Run by hand after vault edits, then commit the regenerated partial and workbook:
`node build.js` stays vault-independent. Not part of the site build.

Usage:
  python .kilo/scripts/gen_ps_audit_page.py [--vault <path>] [--check]

Exit codes: 0 = ok, 1 = drift in --check, 2 = missing vault/note or parse error.
"""
from __future__ import annotations

import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

HERE = Path(__file__).resolve().parents[2]
DEFAULT_VAULT = Path(r"C:\Users\evgeniy.bogdanov\Documents\Personal Vault")
NOTE_REL = Path("Аудит ПС") / "Экспресс-чек-лист-аудита-ПС.md"
PARTIAL = HERE / "site" / "src" / "content" / "genome" / "audit.html"
WORKBOOK = HERE / "site" / "materials" / "ps-express-audit" / "express-audit-ps.xlsx"
DOWNLOAD_RE = re.compile(r"ps-express-audit/express-audit-ps\.xlsx\?v=([^\"'&]+)")

MARKERS = ["META", "PURPOSE", "HOWTO", "COMPOSITION", "SCALE", "CHECKLIST", "SUMMARY"]
CRITERIA_COUNT = 40
BLOCK_COUNT = 6
FULL_CRITERIA = 120

MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня",
          "июля", "августа", "сентября", "октября", "ноября", "декабря"]


class ParseError(Exception):
    pass


# ---------------------------------------------------------------- parsing

def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_separator(cells: list[str]) -> bool:
    return bool(cells) and set("".join(cells)) <= set("-: ")


def read_note(path: Path) -> tuple[dict, str]:
    if not path.exists():
        raise ParseError(f"note not found: {path}")
    text = path.read_text(encoding="utf-8").replace("\r\n", "\n")
    if not text.startswith("---\n"):
        raise ParseError("note has no YAML frontmatter")
    end = text.find("\n---\n", 4)
    if end < 0:
        raise ParseError("frontmatter is not closed")
    fm: dict[str, str] = {}
    for line in text[4:end].split("\n"):
        if ":" in line:
            key, value = line.split(":", 1)
            fm[key.strip()] = value.strip().strip('"').strip("'")
    if not fm.get("version"):
        raise ParseError("frontmatter is missing `version`")
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", fm.get("updated", "")):
        raise ParseError("frontmatter is missing an ISO `updated` date")
    return fm, text[end + 5:]


def split_sections(body: str) -> dict[str, list[str]]:
    sections: dict[str, list[str]] = {}
    current: str | None = None
    buf: list[str] = []
    for line in body.split("\n"):
        m = re.match(r"^##\s+(.+?)\s*$", line)
        if m:
            if current is not None:
                sections[current] = buf
            current = m.group(1)
            buf = []
        elif current is not None:
            buf.append(line)
    if current is not None:
        sections[current] = buf
    return sections


def parse_criteria(lines: list[str]) -> list[dict]:
    blocks: list[dict] = []
    current: dict | None = None
    for line in lines:
        m = re.match(r"^###\s+Блок\s+(\d+)\s+—\s+(.+?)\s*$", line)
        if m:
            current = {"num": int(m.group(1)), "title": m.group(2), "rows": []}
            blocks.append(current)
            continue
        if current is None or not line.lstrip().startswith("|"):
            continue
        cells = split_row(line)
        if not cells or is_separator(cells) or cells[0] == "ID":
            continue
        current["rows"].append(cells)
    return blocks


def parse_composition(lines: list[str]) -> list[list[str]]:
    rows: list[list[str]] = []
    for line in lines:
        s = line.strip()
        if not s.startswith("|"):
            continue
        cells = split_row(s)
        if not is_separator(cells) and cells[0] != "Блок":
            rows.append(cells)
    return rows


def validate_bundle(fm: dict, sections: dict, blocks: list[dict]) -> None:
    required = ["Назначение", "Как пользоваться", "Состав выборки", "Шкала", "Критерии", "Сводка"]
    for name in required:
        if name not in sections:
            raise ParseError(f"section «{name}» is missing")
    if len(blocks) != BLOCK_COUNT:
        raise ParseError(f"expected {BLOCK_COUNT} block headings, found {len(blocks)}")
    if [b["num"] for b in blocks] != list(range(1, BLOCK_COUNT + 1)):
        raise ParseError("block headings must be numbered 1..6 in order")

    seen: list[str] = []
    for block in blocks:
        if not block["rows"]:
            raise ParseError(f"block {block['num']} has no criteria")
        for cells in block["rows"]:
            if len(cells) != 10:
                raise ParseError(f"criterion row has {len(cells)} cells: {cells[0][:40]}")
            m = re.match(r"^(C\d{2})\s*\((\d+\.\d+)\)$", cells[0])
            if not m:
                raise ParseError(f"bad ID cell: {cells[0]!r}")
            cid = m.group(1)
            if int(m.group(2).split(".")[0]) != block["num"]:
                raise ParseError(f"{cid} maps outside block {block['num']}")
            seen.append(cid)
            for index, cell in enumerate(cells[1:], start=1):
                if not cell:
                    raise ParseError(f"{cid}: empty cell {index}")
    expected = [f"C{i:02d}" for i in range(1, CRITERIA_COUNT + 1)]
    if seen != expected:
        raise ParseError("criteria must run C01..C40 sequentially")

    rows = parse_composition(sections["Состав выборки"])
    if len(rows) != BLOCK_COUNT + 1:
        raise ParseError(f"composition table must hold {BLOCK_COUNT} blocks plus a total row")
    try:
        full_total = int(rows[-1][1].replace("*", ""))
        express_total = int(rows[-1][2].replace("*", ""))
    except (IndexError, ValueError) as exc:
        raise ParseError("composition total row is malformed") from exc
    if full_total != FULL_CRITERIA or express_total != CRITERIA_COUNT:
        raise ParseError("composition totals must be 120 and 40")

    scale = [split_row(l) for l in sections["Шкала"] if l.strip().startswith("|")]
    scale = [r for r in scale if not is_separator(r) and r[0] != "Уровень"]
    if len(scale) != 6:
        raise ParseError("scale table must hold six levels")
    if not fm.get("version"):
        raise ParseError("frontmatter is missing `version`")


# ---------------------------------------------------------------- HTML

def esc(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def inline(text: str) -> str:
    text = esc(text)
    text = re.sub(r"\[\[[^\]|]+\|([^\]]+)\]\]", r"\1", text)
    text = re.sub(r"\[\[([^\]]+)\]\]", r"\1", text)
    text = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", text)
    return text


def plain(text: str) -> str:
    text = re.sub(r"\[\[[^\]|]+\|([^\]]+)\]\]", r"\1", text)
    text = re.sub(r"\[\[([^\]]+)\]\]", r"\1", text)
    return text.replace("**", "")


def render_table(rows: list[list[str]], css_class: str = "") -> str:
    if not rows:
        return ""
    head, body = rows[0], rows[1:]
    cls = f' class="table-wrap {css_class}"' if css_class else ' class="table-wrap"'
    out = [f"<div{cls}>", "<table>", "<thead>",
           "<tr>" + "".join(f"<th>{inline(c)}</th>" for c in head) + "</tr>",
           "</thead>", "<tbody>"]
    for row in body:
        total = ' class="is-total"' if row and row[0].replace("*", "").strip().lower() == "итого" else ""
        out.append(f"<tr{total}>" + "".join(f"<td>{inline(c)}</td>" for c in row) + "</tr>")
    out += ["</tbody>", "</table>", "</div>"]
    return "\n".join(out)


def render_flow(lines: list[str]) -> str:
    out: list[str] = []
    i, n = 0, len(lines)
    while i < n:
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                cells = split_row(lines[i])
                if not is_separator(cells):
                    rows.append(cells)
                i += 1
            out.append(render_table(rows))
            continue
        if re.match(r"^\d+[.)]\s+", s) or s.startswith("- "):
            ordered = bool(re.match(r"^\d+[.)]\s+", s))
            items = []
            while i < n:
                cur = lines[i].strip()
                match = re.match(r"^\d+[.)]\s+", cur) if ordered else (cur.startswith("- ") and cur)
                if not match:
                    break
                items.append(re.sub(r"^(\d+[.)]|-)\s+", "", cur))
                i += 1
            tag = "ol" if ordered else "ul"
            out.append(f"<{tag}>" + "".join(f"<li>{inline(t)}</li>" for t in items) + f"</{tag}>")
            continue
        para = []
        while i < n:
            cur = lines[i].strip()
            if not cur or cur.startswith("|") or re.match(r"^(\d+[.)]|-)\s+", cur):
                break
            para.append(cur)
            i += 1
        if para:
            out.append(f"<p>{inline(' '.join(para))}</p>")
    return "\n".join(out)


def ru_date(iso: str) -> str:
    d = datetime.strptime(iso, "%Y-%m-%d")
    return f"{d.day} {MONTHS[d.month - 1]} {d.year}"


def render_meta(fm: dict) -> str:
    return f"Версия {esc(fm['version'])} · обновлено {ru_date(fm['updated'])}"


def render_checklist(blocks: list[dict]) -> str:
    out = ['<div class="audit-chips" aria-label="Блоки критериев">']
    for block in blocks:
        out.append(f'<a href="#block-{block["num"]}">B{block["num"]} · {inline(block["title"])}</a>')
    out.append("</div>")
    out.append('<div class="audit-controls">'
               '<button type="button" class="btn btn-sm" data-audit-toggle="open">Развернуть все</button>'
               '<button type="button" class="btn btn-sm" data-audit-toggle="close">Свернуть все</button>'
               "</div>")
    for block in blocks:
        out.append(f'<section class="audit-block" id="block-{block["num"]}">')
        out.append(f'<h3 class="audit-block-title">Блок {block["num"]} — {inline(block["title"])}</h3>')
        for cells in block["rows"]:
            match = re.match(r"^(C\d{2})\s*\((\d+\.\d+)\)$", cells[0])
            cid, num = match.group(1), match.group(2)
            anchors = cells[2:8]
            out.append(f'<details class="audit-criterion" id="{cid.lower()}">')
            out.append(f'<summary><span class="audit-id">{cid}</span>'
                       f'<span class="audit-num">({num})</span>'
                       f'<span class="audit-title">{inline(cells[1])}</span></summary>')
            out.append('<div class="audit-body">')
            out.append('<div class="table-wrap"><table class="audit-levels">')
            out.append("<thead><tr><th>Уровень</th><th>Что наблюдается</th></tr></thead>")
            out.append("<tbody>")
            for level, text in enumerate(anchors):
                cls = ' class="is-target"' if level == 3 else ""
                out.append(f'<tr{cls}><td class="audit-level">{level}</td><td>{inline(text)}</td></tr>')
            out.append("</tbody></table></div>")
            out.append(f'<p class="audit-kpi"><strong>KPI/результат:</strong> {inline(cells[8])}</p>')
            out.append(f'<p class="audit-verify"><strong>Способ верификации:</strong> {inline(cells[9])}</p>')
            out.append("</div></details>")
        out.append("</section>")
    return "\n".join(out)


def render_regions(fm: dict, sections: dict, blocks: list[dict]) -> dict[str, str]:
    return {
        "META": render_meta(fm),
        "PURPOSE": render_flow(sections["Назначение"]),
        "HOWTO": render_flow(sections["Как пользоваться"]),
        "COMPOSITION": render_flow(sections["Состав выборки"]),
        "SCALE": render_flow(sections["Шкала"]),
        "CHECKLIST": render_checklist(blocks),
        "SUMMARY": render_flow(sections["Сводка"]),
    }


def fill_markers(html: str, regions: dict[str, str]) -> str:
    for tag in MARKERS:
        start, end = f"<!--AUDIT:{tag}:START-->", f"<!--AUDIT:{tag}:END-->"
        if html.count(start) != 1 or html.count(end) != 1:
            raise ParseError(f"marker {tag} must appear exactly once "
                             f"(found {html.count(start)} start / {html.count(end)} end)")
        pattern = re.compile(re.escape(start) + r".*?" + re.escape(end), re.S)
        html = pattern.sub(lambda m, t=tag: start + "\n" + regions[t] + "\n" + end, html, count=1)
    return html


def marker_regions(html: str) -> dict[str, str]:
    found: dict[str, str] = {}
    for tag in MARKERS:
        start, end = f"<!--AUDIT:{tag}:START-->", f"<!--AUDIT:{tag}:END-->"
        pattern = re.compile(re.escape(start) + r"(.*?)" + re.escape(end), re.S)
        match = pattern.search(html)
        if not match:
            raise ParseError(f"marker {tag} is missing")
        found[tag] = match.group(1).strip("\n")
    return found


# ---------------------------------------------------------------- XLSX

def build_workbook(fm: dict, sections: dict, blocks: list[dict]):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.formatting.rule import FormulaRule
    except ImportError as exc:  # pragma: no cover - environment guard
        raise ParseError("openpyxl is required: pip install openpyxl") from exc

    accent = "DC2626"
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=accent)
    bold = Font(bold=True)
    thin = Side(style="thin", color="C8C8C8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap_top = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    ws_instr = wb.active
    ws_instr.title = "Инструкция"
    ws_check = wb.create_sheet("Чек-лист")
    ws_sum = wb.create_sheet("Сводка")
    ws_ref = wb.create_sheet("Справочники")

    updated = datetime.strptime(fm["updated"], "%Y-%m-%d")
    wb.properties.title = "Экспресс-чек-лист аудита производственной системы"
    wb.properties.creator = "AntiManager"
    wb.properties.description = "40 критериев, шкала 0–5, автоматическая сводка"
    wb.properties.created = updated
    wb.properties.modified = updated
    wb.calculation.fullCalcOnLoad = True

    # --- Чек-лист ---------------------------------------------------------
    headers = ["ID", "Блок", "№ в полном листе", "Критерий", "0", "1", "2", "3", "4", "5",
               "KPI/результат", "Способ верификации", "Оценка",
               "Доказательство/комментарий", "Отставание от 3", "Ранг разрыва",
               "Сила", "Ранг силы"]
    for col, name in enumerate(headers, start=1):
        cell = ws_check.cell(row=1, column=col, value=name)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border

    row = 2
    for block in blocks:
        for cells in block["rows"]:
            match = re.match(r"^(C\d{2})\s*\((\d+\.\d+)\)$", cells[0])
            cid, full_ref = match.group(1), match.group(2)
            ws_check.cell(row=row, column=1, value=cid).font = bold
            ws_check.cell(row=row, column=2, value=f"B{block['num']}")
            ref = ws_check.cell(row=row, column=3, value=full_ref)
            ref.number_format = "@"
            ws_check.cell(row=row, column=4, value=cells[1])
            for offset, text in enumerate(cells[2:8]):
                ws_check.cell(row=row, column=5 + offset, value=text)
            ws_check.cell(row=row, column=11, value=cells[8])
            ws_check.cell(row=row, column=12, value=cells[9])
            ws_check.cell(row=row, column=15, value=f'=IF(ISNUMBER(M{row}),3-M{row},"")')
            ws_check.cell(row=row, column=16,
                          value=f'=IF(AND(ISNUMBER(O{row}),O{row}>0),'
                                f'RANK(O{row},$O$2:$O$41,0)+COUNTIF($O$2:O{row},O{row})-1,"")')
            ws_check.cell(row=row, column=17, value=f'=IF(AND(ISNUMBER(M{row}),M{row}>=4),M{row},"")')
            ws_check.cell(row=row, column=18,
                          value=f'=IF(ISNUMBER(Q{row}),'
                                f'RANK(Q{row},$Q$2:$Q$41,0)+COUNTIF($Q$2:Q{row},Q{row})-1,"")')
            for col in range(1, len(headers) + 1):
                cell = ws_check.cell(row=row, column=col)
                cell.border = border
                if col in (4, 5, 6, 7, 8, 9, 10, 11, 12, 14):
                    cell.alignment = wrap_top
            row += 1

    widths = {"A": 7, "B": 7, "C": 9, "D": 46, "E": 22, "F": 22, "G": 22, "H": 22,
              "I": 22, "J": 22, "K": 30, "L": 32, "M": 10, "N": 32, "O": 11,
              "P": 11, "Q": 8, "R": 10}
    for col, width in widths.items():
        ws_check.column_dimensions[col].width = width

    ws_check.freeze_panes = "E2"
    ws_check.auto_filter.ref = "A1:R41"
    ws_check.row_dimensions[1].height = 32
    ws_check.page_setup.orientation = "landscape"
    ws_check.page_setup.paperSize = ws_check.PAPERSIZE_A4
    ws_check.page_setup.fitToWidth = 1
    ws_check.page_setup.fitToHeight = 0
    ws_check.sheet_properties.pageSetUpPr.fitToPage = True
    ws_check.print_title_rows = "1:1"

    dv = DataValidation(type="list", formula1="'Справочники'!$A$2:$A$9", allow_blank=True)
    dv.showErrorMessage = True
    dv.showInputMessage = True
    dv.errorTitle = "Недопустимая оценка"
    dv.error = "Выберите значение из списка: 0–5, «?» или «Н/П»."
    dv.promptTitle = "Оценка"
    dv.prompt = "0–5 по якорям; «?» — нет данных; «Н/П» — не применимо."
    ws_check.add_data_validation(dv)
    dv.add("M2:M41")

    red_fill = PatternFill("solid", fgColor=accent)
    ws_check.conditional_formatting.add(
        "M2:M41", FormulaRule(formula=["AND(ISNUMBER(M2),M2<=2)"], fill=red_fill,
                              font=Font(color="FFFFFF", bold=True)))
    ws_check.conditional_formatting.add(
        "M2:M41", FormulaRule(formula=["AND(ISNUMBER(M2),M2>=4)"], font=Font(bold=True)))

    # --- Справочники ------------------------------------------------------
    ws_ref.cell(row=1, column=1, value="Оценка").font = bold
    for index, value in enumerate([0, 1, 2, 3, 4, 5, "?", "Н/П"], start=2):
        ws_ref.cell(row=index, column=1, value=value)
    ws_ref.column_dimensions["A"].width = 14
    ws_ref.sheet_state = "hidden"

    # --- Сводка -----------------------------------------------------------
    ws_sum.cell(row=1, column=1, value="Сводка экспресс-аудита производственной системы").font = Font(bold=True, size=14)
    ws_sum.cell(row=2, column=1,
                value=f"Источник: экспресс-чек-лист, версия {fm['version']} от {updated.strftime('%d.%m.%Y')}. "
                      "Формулы пересчитываются при открытии файла.")
    ws_sum.merge_cells("A2:H2")
    ws_sum.cell(row=2, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws_sum.row_dimensions[2].height = 32

    sum_head = ["Блок", "Название блока", "Критериев", "Н/П", "Оценено", "Балл блока", "Покрытие"]
    for col, name in enumerate(sum_head, start=1):
        cell = ws_sum.cell(row=4, column=col, value=name)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")

    for index, block in enumerate(blocks):
        r = 5 + index
        label = f"B{block['num']}"
        ws_sum.cell(row=r, column=1, value=label).font = bold
        ws_sum.cell(row=r, column=2, value=block["title"]).alignment = wrap_top
        ws_sum.cell(row=r, column=3, value=f"=COUNTIF('Чек-лист'!$B$2:$B$41,$A{r})")
        ws_sum.cell(row=r, column=4,
                    value=f"=COUNTIFS('Чек-лист'!$B$2:$B$41,$A{r},'Чек-лист'!$M$2:$M$41,\"Н/П\")")
        ws_sum.cell(row=r, column=5,
                    value=f"=COUNTIFS('Чек-лист'!$B$2:$B$41,$A{r},'Чек-лист'!$M$2:$M$41,\">=0\","
                          f"'Чек-лист'!$M$2:$M$41,\"<=5\")")
        ws_sum.cell(row=r, column=6,
                    value=f"=IFERROR(AVERAGEIFS('Чек-лист'!$M$2:$M$41,'Чек-лист'!$B$2:$B$41,$A{r},"
                          f"'Чек-лист'!$M$2:$M$41,\">=0\",'Чек-лист'!$M$2:$M$41,\"<=5\"),\"\")")
        ws_sum.cell(row=r, column=7, value=f"=IFERROR(E{r}/(C{r}-D{r}),\"\")")
        ws_sum.cell(row=r, column=7).number_format = "0%"
        ws_sum.cell(row=r, column=6).number_format = "0.0"

    total = 5 + len(blocks)
    ws_sum.cell(row=total, column=1, value="Итого").font = bold
    ws_sum.cell(row=total, column=3, value=f"=SUM(C5:C{total - 1})")
    ws_sum.cell(row=total, column=4, value=f"=SUM(D5:D{total - 1})")
    ws_sum.cell(row=total, column=5, value=f"=SUM(E5:E{total - 1})")
    ws_sum.cell(row=total, column=6,
                value="=IFERROR(AVERAGEIFS('Чек-лист'!$M$2:$M$41,"
                      "'Чек-лист'!$M$2:$M$41,\">=0\",'Чек-лист'!$M$2:$M$41,\"<=5\"),\"\")")
    ws_sum.cell(row=total, column=7, value=f"=IFERROR(E{total}/(C{total}-D{total}),\"\")")
    ws_sum.cell(row=total, column=6).number_format = "0.0"
    ws_sum.cell(row=total, column=7).number_format = "0%"
    for col in range(1, 8):
        ws_sum.cell(row=total, column=col).font = bold

    metric_row = total + 2
    ws_sum.cell(row=metric_row, column=1, value="Итоговые показатели").font = Font(bold=True, size=12)
    metrics = [
        ("Всего критериев", "=COUNTA('Чек-лист'!$A$2:$A$41)", None),
        ("Применимо", f"=C{total}-D{total}", None),
        ("Оценено", f"=E{total}", None),
        ("Итоговый балл (среднее по оценённым)", f"=F{total}", "0.0"),
        ("Покрытие (оценено / применимо)", f"=G{total}", "0%"),
        ("Число «?»", "=SUMPRODUCT(--('Чек-лист'!$M$2:$M$41=\"?\"))", None),
        ("Число «Н/П»", f"=D{total}", None),
    ]
    for offset, (name, formula, fmt) in enumerate(metrics, start=1):
        r = metric_row + offset
        ws_sum.cell(row=r, column=1, value=name)
        cell = ws_sum.cell(row=r, column=2, value=formula)
        if fmt:
            cell.number_format = fmt

    gaps_row = metric_row + len(metrics) + 2
    ws_sum.cell(row=gaps_row, column=1, value="Топ-5 разрывов (отставание от уровня 3)").font = Font(bold=True, size=12)
    for col, name in enumerate(["№", "ID", "Критерий", "Блок"], start=1):
        cell = ws_sum.cell(row=gaps_row + 1, column=col, value=name)
        cell.font = head_font
        cell.fill = head_fill
    for k in range(1, 6):
        r = gaps_row + 1 + k
        ws_sum.cell(row=r, column=1, value=k)
        match = f"MATCH($A{r},'Чек-лист'!$P$2:$P$41,0)"
        ws_sum.cell(row=r, column=2, value=f"=IFERROR(INDEX('Чек-лист'!$A$2:$A$41,{match}),\"\")")
        ws_sum.cell(row=r, column=3,
                    value=f"=IFERROR(INDEX('Чек-лист'!$D$2:$D$41,{match}),\"\")").alignment = wrap_top
        ws_sum.cell(row=r, column=4, value=f"=IFERROR(INDEX('Чек-лист'!$B$2:$B$41,{match}),\"\")")

    strengths_row = gaps_row + 8
    ws_sum.cell(row=strengths_row, column=1, value="Топ-5 сильных сторон (уровни 4–5)").font = Font(bold=True, size=12)
    for col, name in enumerate(["№", "ID", "Критерий", "Блок"], start=1):
        cell = ws_sum.cell(row=strengths_row + 1, column=col, value=name)
        cell.font = head_font
        cell.fill = head_fill
    for k in range(1, 6):
        r = strengths_row + 1 + k
        ws_sum.cell(row=r, column=1, value=k)
        match = f"MATCH($A{r},'Чек-лист'!$R$2:$R$41,0)"
        ws_sum.cell(row=r, column=2, value=f"=IFERROR(INDEX('Чек-лист'!$A$2:$A$41,{match}),\"\")")
        ws_sum.cell(row=r, column=3,
                    value=f"=IFERROR(INDEX('Чек-лист'!$D$2:$D$41,{match}),\"\")").alignment = wrap_top
        ws_sum.cell(row=r, column=4, value=f"=IFERROR(INDEX('Чек-лист'!$B$2:$B$41,{match}),\"\")")

    for col, width in {"A": 44, "B": 16, "C": 60, "D": 8, "E": 12, "F": 13, "G": 12}.items():
        ws_sum.column_dimensions[col].width = width
    ws_sum.cell(row=strengths_row + 7, column=1,
                value="Пустые строки топ-5 означают, что оценок соответствующего вида пока нет.").font = Font(italic=True)

    # --- Инструкция -------------------------------------------------------
    ws_instr.column_dimensions["A"].width = 112
    ws_instr.column_dimensions["B"].width = 20
    cursor = 1

    def line(text, *, bold_=False, size=None, italic=False, wrap=True):
        nonlocal cursor
        cell = ws_instr.cell(row=cursor, column=1, value=text)
        cell.font = Font(bold=bold_, italic=italic, size=size or 11)
        cell.alignment = Alignment(wrap_text=wrap, vertical="top")
        cursor += 1

    line("Экспресс-чек-лист аудита производственной системы", bold_=True, size=16)
    line(f"Версия {fm['version']} · {updated.strftime('%d.%m.%Y')} · AntiManager")
    cursor += 1

    line("Назначение", bold_=True, size=12)
    for para in re.split(r"\n\s*\n", "\n".join(sections["Назначение"]).strip()):
        if para.strip():
            line(plain(para.replace("\n", " ")))
    cursor += 1

    line("Как заполнять", bold_=True, size=12)
    for item in [l.strip()[2:] for l in sections["Как пользоваться"] if l.strip().startswith("- ")]:
        line("• " + plain(item))
    cursor += 1

    line("Шкала", bold_=True, size=12)
    scale_rows = [split_row(l) for l in sections["Шкала"] if l.strip().startswith("|")]
    for cells in scale_rows:
        if is_separator(cells) or cells[0] == "Уровень":
            continue
        line(f"{cells[0]} — {cells[1]}")
    cursor += 1

    line("Правило достоверности", bold_=True, size=12)
    line("Интервью — не выше 2. Запись — 3. Запись и наблюдение — 4. "
         "Запись, данные, наблюдение и измеренный эффект — 5.")
    line("«Н/П» ставится только с письменным обоснованием. «?» ставится при отсутствии данных.")
    line("«?» и «Н/П» не входят в среднее; оценка вводится в столбец M из списка 0–5.")
    cursor += 1

    line("Показатели", bold_=True, size=12)
    line("Балл блока — среднее по оценённым критериям блока.")
    line("Итоговый балл — среднее по всем оценённым критериям листа; критерии равноправны.")
    line("Покрытие — доля оценённых критериев от применимых (критерии минус «Н/П»).")
    line("Топ-5 разрывов — критерии с наибольшим отставанием от уровня 3.")
    line("Топ-5 сильных сторон — критерии с оценкой 4–5.")
    cursor += 1

    line("Пересчёт", bold_=True, size=12)
    line("Формулы пересчитываются при открытии файла. Если значения не появились, нажмите F9.")
    line("Файл не защищён: столбцы можно дополнять, формулы не удаляйте.")

    return wb


def snapshot(wb) -> dict:
    snap = {"sheets": list(wb.sheetnames), "cells": [], "validations": [], "hidden": []}
    for name in ("Инструкция", "Чек-лист", "Сводка", "Справочники"):
        ws = wb[name]
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                if cell.value is not None:
                    snap["cells"].append((name, cell.coordinate, cell.value))
    ws = wb["Чек-лист"]
    for dv in ws.data_validations.dataValidation:
        snap["validations"].append((dv.type, str(dv.formula1), str(dv.sqref)))
    snap["hidden"] = [(name, wb[name].sheet_state) for name in wb.sheetnames]
    return snap


# ---------------------------------------------------------------- main

def generate(vault: Path, check: bool) -> int:
    note = vault / NOTE_REL
    try:
        fm, body = read_note(note)
        sections = split_sections(body)
        blocks = parse_criteria(sections.get("Критерии", []))
        validate_bundle(fm, sections, blocks)
        regions = render_regions(fm, sections, blocks)
        html = PARTIAL.read_text(encoding="utf-8")
        rendered = fill_markers(html, regions)
        wb = build_workbook(fm, sections, blocks)
    except ParseError as exc:
        print(f"ERROR {exc}", file=sys.stderr)
        return 2

    match = DOWNLOAD_RE.search(html)
    if not match:
        print("WARN partial has no express-audit-ps.xlsx download link", file=sys.stderr)
    elif match.group(1) != fm["version"]:
        print(f"WARN download href is ?v={match.group(1)} but the vault version is {fm['version']}",
              file=sys.stderr)

    if check:
        problems = 0
        try:
            actual_regions = marker_regions(html)
        except ParseError as exc:
            print(f"ERROR {exc}", file=sys.stderr)
            return 1
        for tag in MARKERS:
            if actual_regions[tag] != regions[tag]:
                print(f"DRIFT HTML region {tag}", file=sys.stderr)
                problems += 1
        if not WORKBOOK.exists():
            print(f"DRIFT workbook missing: {WORKBOOK}", file=sys.stderr)
            problems += 1
        else:
            try:
                from openpyxl import load_workbook
            except ImportError as exc:  # pragma: no cover
                print(f"ERROR {exc}", file=sys.stderr)
                return 2
            actual = load_workbook(WORKBOOK, data_only=False)
            if snapshot(actual) != snapshot(wb):
                print("DRIFT workbook content", file=sys.stderr)
                problems += 1
        if problems:
            print(f"\n{problems} drift(s) — run the generator without --check", file=sys.stderr)
            return 1
        print("OK no drift")
        return 0

    PARTIAL.write_text(rendered, encoding="utf-8", newline="\n")
    WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
    wb.save(WORKBOOK)
    print(f"wrote {PARTIAL.relative_to(HERE)} ({len(rendered)} chars)")
    print(f"wrote {WORKBOOK.relative_to(HERE)} ({WORKBOOK.stat().st_size} bytes)")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--vault", type=Path, default=DEFAULT_VAULT)
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    if not args.vault.exists():
        print(f"ERROR vault not found: {args.vault}", file=sys.stderr)
        return 2
    return generate(args.vault, args.check)


if __name__ == "__main__":
    sys.exit(main())
